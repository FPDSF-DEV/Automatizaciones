"""Pruebas del Modelo PSS: Excel de afiliación, ámbito, pirámide y los 32 gráficos."""

from pathlib import Path

import pytest

from terridata import graficas as g
from terridata.estilo import CLARO
from terridata.lectura import Fila
from terridata.pss.afiliacion import Afiliacion, leer_cobertura
from terridata.pss.ambito import (
    GRUPOS_QUINQUENALES,
    datos_piramide,
    detectar_ambito,
    municipios_de,
    promedio_ponderado,
    suma,
)
from terridata.pss.catalogo_pss import CATALOGO_PSS, Datos
from terridata.pss.graficos_pss import DEFINICION_ODS, ESPECIFICACIONES, generar

MUNICIPIOS = ["Agrado", "Gigante", "Garzón"]


# --------------------------------------------------------------------------- helpers


def fila(entidad, dimension, indicador, valor=None, anio=2026, texto=None,
         cod="41306", unidad="Personas"):
    return Fila(
        cod_departamento="41", departamento="Huila", cod_entidad=cod, entidad=entidad,
        dimension=dimension, subcategoria="x", indicador=indicador, valor=valor,
        texto=texto, anio=anio, mes=12, fuente="DANE", unidad=unidad,
    )


def filas_demograficas(entidad, hombres, mujeres, urbana, anio=2026, cod="41306"):
    return [
        fila(entidad, "Demografía y población", "Población total de hombres", hombres, anio, cod=cod),
        fila(entidad, "Demografía y población", "Población total de mujeres", mujeres, anio, cod=cod),
        fila(entidad, "Demografía y población", "Población urbana", urbana, anio, cod=cod),
    ]


def filas_quinquenios(entidad, dimension, anio, hombres=100.0, mujeres=110.0, solo_mujeres=False):
    filas = []
    for grupo in GRUPOS_QUINQUENALES:
        if not solo_mujeres:
            filas.append(fila(entidad, dimension, f"Población de hombres de {grupo}", hombres, anio))
        filas.append(fila(entidad, dimension, f"Población de mujeres de {grupo}", mujeres, anio))
    return filas


def construir_datos(filas, anio=2026, cobertura=None):
    from terridata.lectura import Lectura
    from terridata.resolucion import resolver

    entidades = []
    for f in filas:
        if f.entidad not in entidades:
            entidades.append(f.entidad)
    lectura = Lectura(filas=filas, entidades=entidades, departamento="Huila", archivos=[])
    res = resolver(filas, CATALOGO_PSS, anio_corte=anio, entidades=entidades,
                   departamento="Huila")
    ambito = detectar_ambito(entidades, "Huila")
    codigos = {}
    for f in filas:
        codigos.setdefault(f.entidad, f.cod_entidad)
    return Datos(lectura=lectura, res=res, cobertura=cobertura, ambito=ambito,
                 municipios=municipios_de(ambito), anio=anio, codigos=codigos,
                 cod_departamento="41")


# --------------------------------------------------------------------------- Excel


@pytest.fixture
def excel(tmp_path):
    """Un Excel de ADRES en miniatura, con el mismo formato del real."""
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "CoberturaDepartamento"
    hoja.append(["CIFRAS JULIO-2026"])
    hoja.append(["IdDepartamento", "Departamento", "Contributivo", "Subsidiado",
                 "Excepción & Especiales", "Afiliados", "Población", "Cobertura",
                 "Afil. Subsidiado sin SISBEN IV"])
    hoja.append(["00", "SIN UBICACIÓN", None, None, 100, 100, 0, None, 0])
    hoja.append(["41", "HUILA", 327943, 858759, 28963, 1215665, 1222896, 0.994, 95841])

    hoja2 = libro.create_sheet("CoberturaMunicipio")
    hoja2.append(["CIFRAS JULIO-2026"])
    hoja2.append(["IdDepartamento", "Departamento", "IdMunicipio", "Municipio",
                  "Contributivo", "Subsidiado", "Excepción & Especiales", "Afiliados",
                  "Población", "Cobertura", "Afil. Subsidiado sin SISBEN IV"])
    hoja2.append(["41", "HUILA", "41013", "AGRADO", 857, 7639, 174, 8670, 9404, 0.9219, 367])
    hoja2.append(["41", "HUILA", "41306", "GIGANTE", 3620, 19219, 661, 23500, 23337, 1, 1459])

    destino = tmp_path / "CIFRAS JULIO-2026.xlsx"
    libro.save(destino)
    return destino


def test_lee_departamentos_y_municipios(excel):
    c = leer_cobertura(excel)
    assert c.corte == "JULIO-2026"
    assert set(c.departamentos) == {"41"}
    assert set(c.municipios) == {"41013", "41306"}


def test_descarta_la_fila_sin_ubicacion(excel):
    c = leer_cobertura(excel)
    assert "00" not in c.departamentos


def test_busca_las_columnas_por_nombre(excel):
    """El orden de las columnas cambia entre cortes: se ubican por encabezado."""
    c = leer_cobertura(excel)
    huila = c.departamento("41")
    assert huila.contributivo == 327943
    assert huila.subsidiado == 858759
    assert huila.poblacion == 1222896


def test_codigo_de_municipio_con_ceros(excel):
    c = leer_cobertura(excel)
    assert c.municipio("41013") is not None
    assert c.municipio(41013) is not None      # acepta entero


def test_cobertura_en_porcentaje_y_topada():
    a = Afiliacion("41306", "GIGANTE", "HUILA", 1.0, 1.0, 1.0, 3.0, 2.0, 1.2, 0.0)
    assert a.cobertura_pct == 100.0
    assert a.sin_afiliar_pct == 0.0


def test_brecha_es_el_complemento():
    a = Afiliacion("41013", "AGRADO", "HUILA", 1.0, 1.0, 1.0, 3.0, 10.0, 0.9219, 0.0)
    assert a.cobertura_pct == pytest.approx(92.19)
    assert a.sin_afiliar_pct == pytest.approx(7.81)


def test_excel_inexistente():
    with pytest.raises(FileNotFoundError):
        leer_cobertura("no/existe.xlsx")


def test_excel_sin_las_hojas(tmp_path):
    import openpyxl

    libro = openpyxl.Workbook()
    destino = tmp_path / "otro.xlsx"
    libro.save(destino)
    with pytest.raises(ValueError, match="CoberturaDepartamento"):
        leer_cobertura(destino)


# --------------------------------------------------------------------------- ámbito


def test_sin_extraccion_del_departamento_el_ambito_es_agregado():
    a = detectar_ambito(["Gigante", "Garzón"], "Huila")
    assert not a.es_departamento
    assert "agregado" in a.calificador
    assert municipios_de(a) == ["Gigante", "Garzón"]


def test_con_extraccion_del_departamento_se_usa_esa_entidad():
    a = detectar_ambito(["Gigante", "Huila", "Garzón"], "Huila")
    assert a.es_departamento and a.nombre == "Huila"
    assert municipios_de(a) == ["Gigante", "Garzón"]


def test_suma_ignora_los_faltantes():
    assert suma([1.0, None, 3.0]) == 4.0
    assert suma([None, None]) is None


def test_las_tasas_se_promedian_ponderando_por_poblacion():
    """Sumar tasas no significa nada; se ponderan por población."""
    promedio, aportan = promedio_ponderado([10.0, 20.0], [1000.0, 3000.0])
    assert promedio == pytest.approx(17.5)
    assert aportan == 2


def test_promedio_ponderado_ignora_entidades_sin_dato():
    promedio, aportan = promedio_ponderado([10.0, None], [1000.0, 3000.0])
    assert promedio == pytest.approx(10.0) and aportan == 1


def test_promedio_ponderado_sin_datos():
    assert promedio_ponderado([None, None], [1.0, 2.0]) == (None, 0)


# --------------------------------------------------------------------------- pirámide


def test_la_piramide_exige_los_dos_sexos_completos():
    """Con la serie masculina incompleta se retrocede al año que sí tiene ambas."""
    filas = (
        filas_quinquenios("Gigante", "Demografía y población", 2026, solo_mujeres=True)
        + filas_quinquenios("Gigante", "Censo 2005 y proyecciones DANE", 2019,
                            hombres=90.0, mujeres=95.0)
    )
    p = datos_piramide(filas, ["Gigante"], 2026)
    assert p.completa
    assert p.anio == 2019
    assert p.dimension == "Censo 2005 y proyecciones DANE"
    assert p.hombres[0] == 90.0 and p.mujeres[0] == 95.0


def test_la_piramide_usa_el_anio_pedido_si_esta_completo():
    filas = filas_quinquenios("Gigante", "Demografía y población", 2026)
    p = datos_piramide(filas, ["Gigante"], 2026)
    assert p.anio == 2026 and p.completa


def test_la_piramide_suma_las_entidades():
    filas = (filas_quinquenios("Gigante", "Demografía y población", 2026)
             + filas_quinquenios("Garzón", "Demografía y población", 2026))
    p = datos_piramide(filas, ["Gigante", "Garzón"], 2026)
    assert p.hombres[0] == 200.0 and p.mujeres[0] == 220.0


def test_sin_quinquenios_no_hay_piramide():
    p = datos_piramide([fila("Gigante", "Salud", "x", 1.0)], ["Gigante"], 2026)
    assert not p.completa and p.anio is None


# --------------------------------------------------------------------------- datos


def test_la_poblacion_total_es_hombres_mas_mujeres():
    """TerriData no publica el total proyectado: se arma con H + M (verificado
    contra el documento: Gigante 2026 = 11.480 + 11.873 = 23.353)."""
    d = construir_datos(filas_demograficas("Gigante", 11480.0, 11873.0, 9334.0))
    assert d.poblacion("Gigante") == 23353.0


def test_la_rural_se_calcula_si_la_directa_es_de_otro_anio():
    """La 'Población rural' directa solo existe en la serie del censo (2019)."""
    filas = filas_demograficas("Gigante", 11480.0, 11873.0, 9334.0, anio=2026)
    filas.append(fila("Gigante", "Censo 2005 y proyecciones DANE", "Población rural",
                      17505.0, 2019))
    d = construir_datos(filas)
    assert d.urbana("Gigante") == 9334.0
    assert d.rural("Gigante") == pytest.approx(23353.0 - 9334.0)


def test_la_rural_directa_se_usa_si_es_del_mismo_anio():
    filas = filas_demograficas("Gigante", 100.0, 100.0, 60.0, anio=2026)
    filas.append(fila("Gigante", "Demografía y población", "Población rural", 140.0, 2026))
    d = construir_datos(filas)
    assert d.rural("Gigante") == 140.0


def test_cabecera_mas_rural_es_la_poblacion_total():
    d = construir_datos(filas_demograficas("Gigante", 11480.0, 11873.0, 9334.0))
    assert d.urbana("Gigante") + d.rural("Gigante") == d.poblacion("Gigante")


def test_valor_retrocede_al_anio_anterior_con_dato():
    filas = [fila("Gigante", "Salud", "Cobertura vacunación pentavalente en menores de 1 año",
                  50.0, 2023, unidad="Porcentaje")]
    d = construir_datos(filas)
    valor, anio = d.valor("VAC_PENTA", "Gigante", 2025)
    assert (valor, anio) == (50.0, 2023)


def test_valor_sin_dato_anterior_es_none():
    filas = [fila("Gigante", "Salud", "Cobertura vacunación pentavalente en menores de 1 año",
                  50.0, 2025, unidad="Porcentaje")]
    d = construir_datos(filas)
    assert d.valor("VAC_PENTA", "Gigante", 2018) == (None, None)


# --------------------------------------------------------------------------- catálogo


def test_son_exactamente_32_graficos():
    assert len(ESPECIFICACIONES) == 32
    assert [s.numero for s in ESPECIFICACIONES] == list(range(1, 33))


def test_los_ids_no_se_repiten():
    ids = [s.id for s in ESPECIFICACIONES]
    assert len(ids) == len(set(ids))


def test_los_17_ods_estan_completos():
    assert len(DEFINICION_ODS) == 17
    numeros = [d[0] for d in DEFINICION_ODS]
    assert numeros == list(range(16, 33))


def test_el_nombre_de_archivo_lleva_el_numero():
    from terridata.pss.catalogo_pss import Render

    render = Render(ESPECIFICACIONES[0], None, "t", "s", "p")
    assert render.archivo == "g01_piramide_poblacional"


def test_todos_los_indicadores_ods_estan_en_el_catalogo():
    ids = {i.id for i in CATALOGO_PSS}
    for _n, _id, _t, indicador, _a, _l, _nota in DEFINICION_ODS:
        assert indicador in ids, indicador


def test_un_grafico_sin_datos_se_omite_sin_excepcion():
    d = construir_datos([fila("Gigante", "Salud", "x", 1.0)])
    omitidos = [spec for spec, render in generar(d, CLARO) if render is None]
    assert len(omitidos) == 32          # no hay datos de nada, pero nada explota


def test_los_graficos_de_aseguramiento_se_omiten_sin_excel():
    d = construir_datos(filas_demograficas("Gigante", 100.0, 100.0, 60.0))
    resultados = dict((spec.numero, render) for spec, render in generar(d, CLARO, [9, 10, 11, 12]))
    assert all(r is None for r in resultados.values())


def test_se_arman_los_graficos_demograficos_con_datos_minimos():
    filas = []
    for nombre, cod in (("Gigante", "41306"), ("Garzón", "41298")):
        filas += filas_demograficas(nombre, 100.0, 110.0, 60.0, cod=cod)
        filas += filas_quinquenios(nombre, "Demografía y población", 2026)
    d = construir_datos(filas)
    hechos = {spec.numero: render for spec, render in generar(d, CLARO, [1, 2, 3, 13, 14, 15])}
    for numero, render in hechos.items():
        assert render is not None, numero
        assert render.encabezados and render.filas
        g.cerrar(render.fig)


# --------------------------------------------------------------------------- primitivas nuevas


def test_la_piramide_dibuja_los_dos_lados():
    fig = g.piramide(["0-4", "5-9"], [10.0, 20.0], [12.0, 18.0], CLARO, "Título")
    ax = fig.axes[0]
    extremos = [p.get_path().get_extents() for p in ax.patches]
    assert any(e.x0 < 0 for e in extremos)     # hombres a la izquierda
    assert any(e.x1 > 0 for e in extremos)     # mujeres a la derecha
    assert ax.get_xlim()[0] == -ax.get_xlim()[1]
    g.cerrar(fig)


def test_barras_agrupadas_h_mantiene_el_orden_de_las_series():
    """Con el eje invertido, la primera serie debe quedar arriba (regresión)."""
    fig = g.barras_agrupadas_h(
        ["A", "B"], [("Primera", [1.0, 2.0]), ("Segunda", [3.0, 4.0])], CLARO, "Título")
    ax = fig.axes[0]
    centros = [
        (p.get_path().get_extents().y0 + p.get_path().get_extents().y1) / 2
        for p in ax.patches
    ]
    # cuatro barras: dos de la primera serie y dos de la segunda, alternadas por grupo
    primera = sorted(centros)[:2]
    assert all(c < 0.5 or (0.5 < c < 1.5) for c in primera)
    # dentro del primer grupo, la primera serie tiene menor y (arriba con eje invertido)
    del_grupo = sorted(c for c in centros if c < 0.5)
    assert len(del_grupo) == 2 and del_grupo[0] < del_grupo[1]
    g.cerrar(fig)


def test_ranking_descendente_pone_el_mayor_arriba():
    fig = g.barras_ranking(["A", "B", "C"], [1.0, 3.0, 2.0], CLARO, "T", orden="desc")
    ax = fig.axes[0]
    assert [t.get_text() for t in ax.get_yticklabels()] == ["B", "C", "A"]
    g.cerrar(fig)


def test_ranking_en_orden_dado_no_reordena():
    fig = g.barras_ranking(["Hombres", "Mujeres"], [10.0, 5.0], CLARO, "T", orden="dado")
    ax = fig.axes[0]
    assert [t.get_text() for t in ax.get_yticklabels()] == ["Hombres", "Mujeres"]
    g.cerrar(fig)


def test_etiquetas_de_valor_personalizadas():
    fig = g.barras_ranking(["Hombres", "Mujeres"], [10.0, 5.0], CLARO, "T", orden="dado",
                           etiquetas_valor={"Hombres": "10 · 66,7 %"})
    textos = [t.get_text() for t in fig.axes[0].texts]
    assert "10 · 66,7 %" in textos
    g.cerrar(fig)


def test_la_dona_deja_el_centro_hueco():
    fig = g.torta(["A", "B"], [60.0, 40.0], CLARO, "T", dona=True, centro="429\nx 100.000")
    textos = [t.get_text() for t in fig.axes[0].texts]
    assert "429" in textos
    g.cerrar(fig)
