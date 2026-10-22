"""Lector del Excel de afiliación al SGSSS (ADRES — "CIFRAS <MES>-<AÑO>.xlsx").

TerriData no publica las cifras de afiliación de corte, así que los gráficos de
aseguramiento del Modelo PSS se alimentan de este archivo. Se leen dos hojas:

* `CoberturaDepartamento` — una fila por departamento.
* `CoberturaMunicipio` — una fila por municipio, con el código DANE de 5 dígitos,
  que es la llave para cruzar con TerriData.

El archivo trae dos filas de preámbulo (el título "CIFRAS ..." y el encabezado),
así que las columnas se localizan **por nombre**, no por posición: el orden
cambia entre cortes mensuales.
"""

from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger("terridata.pss.afiliacion")

HOJA_DEPARTAMENTO = "CoberturaDepartamento"
HOJA_MUNICIPIO = "CoberturaMunicipio"

# nombre lógico -> cómo aparece el encabezado en la hoja (ya normalizado)
COLUMNAS = {
    "cod_departamento": "iddepartamento",
    "departamento": "departamento",
    "cod_municipio": "idmunicipio",
    "municipio": "municipio",
    "contributivo": "contributivo",
    "subsidiado": "subsidiado",
    "excepcion": "excepcion & especiales",
    "afiliados": "afiliados",
    "poblacion": "poblacion",
    "cobertura": "cobertura",
    "subsidiado_sin_sisben": "afil. subsidiado sin sisben iv",
}


@dataclass(frozen=True)
class Afiliacion:
    codigo: str
    nombre: str
    departamento: str
    contributivo: float | None
    subsidiado: float | None
    excepcion: float | None
    afiliados: float | None
    poblacion: float | None
    cobertura: float | None          # 0-1 tal como viene en el archivo
    subsidiado_sin_sisben: float | None

    @property
    def cobertura_pct(self) -> float | None:
        """Cobertura en porcentaje, topada en 100 (el archivo trae 1,0 = 100 %)."""
        if self.cobertura is None:
            return None
        return min(self.cobertura * 100.0, 100.0)

    @property
    def sin_afiliar_pct(self) -> float | None:
        c = self.cobertura_pct
        return None if c is None else max(100.0 - c, 0.0)


@dataclass
class Cobertura:
    """Lo que trae el archivo, listo para consultar por código DANE."""

    archivo: Path
    corte: str                                   # "JULIO-2026" según el título
    departamentos: dict[str, Afiliacion]         # por código de 2 dígitos
    municipios: dict[str, Afiliacion]            # por código de 5 dígitos

    def departamento(self, codigo: str) -> Afiliacion | None:
        return self.departamentos.get(str(codigo).zfill(2))

    def municipio(self, codigo: str) -> Afiliacion | None:
        return self.municipios.get(str(codigo).zfill(5))


def _normalizar(texto) -> str:
    """Minúsculas sin tildes, para comparar encabezados sin depender de la ortografía."""
    if texto is None:
        return ""
    sin_tildes = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return " ".join(sin_tildes.lower().split())


def _a_numero(valor) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _indice_columnas(encabezado: list) -> dict[str, int]:
    """Ubica cada columna por su nombre normalizado."""
    normalizados = [_normalizar(c) for c in encabezado]
    indices: dict[str, int] = {}
    for logico, buscado in COLUMNAS.items():
        for i, nombre in enumerate(normalizados):
            if nombre == buscado or nombre.startswith(buscado):
                indices[logico] = i
                break
    return indices


def _fila_encabezado(filas: list) -> int:
    """La hoja arranca con el título del corte; el encabezado es la fila siguiente."""
    for i, fila in enumerate(filas[:6]):
        if any(_normalizar(c) == "iddepartamento" for c in fila):
            return i
    return 1


def _leer_hoja(ws, es_municipio: bool) -> tuple[list[Afiliacion], str]:
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    if not filas:
        return [], ""
    titulo = next((str(c) for c in filas[0] if c), "")
    inicio = _fila_encabezado(filas)
    indices = _indice_columnas(filas[inicio])

    faltan = [
        c for c in ("cod_departamento", "departamento", "afiliados", "poblacion")
        if c not in indices
    ]
    if faltan:
        raise ValueError(
            f"la hoja '{ws.title}' no tiene las columnas {faltan}; "
            "¿cambió el formato del archivo de ADRES?"
        )

    registros: list[Afiliacion] = []
    for fila in filas[inicio + 1:]:
        def col(nombre):
            i = indices.get(nombre)
            return fila[i] if i is not None and i < len(fila) else None

        codigo = col("cod_municipio") if es_municipio else col("cod_departamento")
        if codigo is None or str(codigo).strip() == "":
            continue
        codigo = str(codigo).strip().zfill(5 if es_municipio else 2)
        if not codigo.isdigit() or codigo.startswith("00"):
            continue      # la fila "SIN UBICACIÓN" no es un territorio

        nombre = col("municipio") if es_municipio else col("departamento")
        registros.append(
            Afiliacion(
                codigo=codigo,
                nombre=str(nombre or "").strip(),
                departamento=str(col("departamento") or "").strip(),
                contributivo=_a_numero(col("contributivo")),
                subsidiado=_a_numero(col("subsidiado")),
                excepcion=_a_numero(col("excepcion")),
                afiliados=_a_numero(col("afiliados")),
                poblacion=_a_numero(col("poblacion")),
                cobertura=_a_numero(col("cobertura")),
                subsidiado_sin_sisben=_a_numero(col("subsidiado_sin_sisben")),
            )
        )
    return registros, titulo


def leer_cobertura(ruta: str | Path) -> Cobertura:
    """Lee el Excel de ADRES. Lanza si el archivo no existe o le faltan columnas."""
    import openpyxl

    ruta = Path(ruta)
    if not ruta.is_file():
        raise FileNotFoundError(f"no encuentro el Excel de afiliación: {ruta}")

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        for hoja in (HOJA_DEPARTAMENTO, HOJA_MUNICIPIO):
            if hoja not in libro.sheetnames:
                raise ValueError(
                    f"'{ruta.name}' no tiene la hoja '{hoja}'. "
                    f"Hojas encontradas: {', '.join(libro.sheetnames)}"
                )
        deps, titulo = _leer_hoja(libro[HOJA_DEPARTAMENTO], es_municipio=False)
        muns, _ = _leer_hoja(libro[HOJA_MUNICIPIO], es_municipio=True)
    finally:
        libro.close()

    corte = ""
    if titulo:
        encontrado = re.search(r"CIFRAS\s+(.+)", titulo, re.IGNORECASE)
        corte = encontrado.group(1).strip() if encontrado else titulo.strip()

    log.info("afiliación: %d departamentos y %d municipios (corte %s)",
             len(deps), len(muns), corte or "sin identificar")

    return Cobertura(
        archivo=ruta,
        corte=corte,
        departamentos={d.codigo: d for d in deps},
        municipios={m.codigo: m for m in muns},
    )


def buscar_excel(carpeta: str | Path = ".") -> Path | None:
    """El .xlsx de ADRES más reciente en la carpeta, si hay alguno."""
    candidatos = [
        p for p in Path(carpeta).glob("*.xlsx")
        if not p.name.startswith("~$") and "cifras" in p.name.lower()
    ]
    if not candidatos:
        candidatos = [p for p in Path(carpeta).glob("*.xlsx") if not p.name.startswith("~$")]
    if not candidatos:
        return None
    return max(candidatos, key=lambda p: p.stat().st_mtime)
