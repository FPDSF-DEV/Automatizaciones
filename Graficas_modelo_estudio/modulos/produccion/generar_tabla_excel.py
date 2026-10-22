from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import DefaultDict, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


YEAR_PATTERN = re.compile(r"(\d{4})")
# Misma fuente que generator.py: la carpeta SIHO con estructura de SJDC_SIHO.
SIHO_PATTERN = "**/Formularios_Anuales/*/Produccion/*.csv"
PROJECT_ROOT = Path(__file__).resolve().parents[2]


def extract_year_from_filename(file_path: Path) -> int:
	match = YEAR_PATTERN.search(file_path.stem)
	if not match:
		raise ValueError(f"No se pudo detectar el anio en el archivo: {file_path.name}")
	return int(match.group(1))


def normalize_concept_name(value: str) -> str:
	text = value.strip()
	text = re.sub(r"^[\"'“”‘’]+|[\"'“”‘’]+$", "", text)
	return text.strip()


def parse_numeric_value(raw_value: str) -> float | None:
	raw_value = raw_value.strip()
	if not raw_value:
		return None
	try:
		return float(raw_value)
	except ValueError:
		return None


def load_yearly_totals(
	input_dir: Path,
	concept_column: str,
	value_column: str,
) -> Dict[int, Dict[str, float]]:
	csv_files = sorted(input_dir.glob(SIHO_PATTERN), key=extract_year_from_filename)
	if not csv_files:
		raise FileNotFoundError(f"No se encontraron archivos CSV ({SIHO_PATTERN}) en: {input_dir}")

	yearly_data: Dict[int, DefaultDict[str, float]] = {}

	for csv_file in csv_files:
		year = extract_year_from_filename(csv_file)
		if year in yearly_data:
			raise ValueError(f"Hay dos reportes de Produccion para {year}; deje solo uno por anio: {csv_file}")
		year_totals: DefaultDict[str, float] = defaultdict(float)

		with csv_file.open("r", encoding="utf-8-sig", newline="") as handle:
			reader = csv.reader(handle)
			try:
				header = next(reader)
			except StopIteration:
				continue

			header = [column.strip() for column in header]

			missing_columns = [column for column in (concept_column, value_column) if column not in header]
			if missing_columns:
				missing_text = ", ".join(missing_columns)
				raise ValueError(f"{csv_file.name} no tiene las columnas requeridas: {missing_text}")

			concept_idx = header.index(concept_column)
			value_idx = header.index(value_column)
			expected_len = len(header)
			fixed_tail_size = expected_len - value_idx - 1

			for raw_row in reader:
				if not raw_row:
					continue
				if len(raw_row) < expected_len:
					continue

				if len(raw_row) == expected_len:
					concept = normalize_concept_name(raw_row[concept_idx])
					value_raw = raw_row[value_idx]
				else:
					extra_fields = len(raw_row) - expected_len
					concept_end = concept_idx + 1 + extra_fields
					concept = normalize_concept_name(",".join(raw_row[concept_idx:concept_end]))
					value_raw = raw_row[value_idx + extra_fields]

				if not concept:
					continue

				value = parse_numeric_value(value_raw)
				if value is None:
					continue

				year_totals[concept] += value

		yearly_data[year] = year_totals

	return yearly_data


def build_excel_rows(yearly_data: Dict[int, Dict[str, float]]) -> List[List[object]]:
	if not yearly_data:
		return []

	years = sorted(yearly_data.keys())
	all_concepts: List[str] = []
	seen_concepts: set[str] = set()

	for year in years:
		for concept in yearly_data[year].keys():
			if concept in seen_concepts:
				continue
			seen_concepts.add(concept)
			all_concepts.append(concept)

	rows: List[List[object]] = [["concepto", *[str(year) for year in years]]]

	for concept in all_concepts:
		row: List[object] = [concept]
		for year in years:
			row.append(yearly_data[year].get(concept))
		rows.append(row)

	return rows


def write_excel(output_file: Path, sheet_name: str, rows: List[List[object]]) -> None:
	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = sheet_name[:31] if sheet_name else "tabla"

	for row in rows:
		worksheet.append(row)

	if worksheet.max_row >= 1:
		for cell in worksheet[1]:
			cell.font = Font(bold=True)

	worksheet.freeze_panes = "B2"
	worksheet.auto_filter.ref = worksheet.dimensions

	for column_cells in worksheet.columns:
		max_length = 0
		column_letter = get_column_letter(column_cells[0].column)
		for cell in column_cells:
			cell_value = "" if cell.value is None else str(cell.value)
			max_length = max(max_length, len(cell_value))
		worksheet.column_dimensions[column_letter].width = min(max_length + 2, 80)

	output_file.parent.mkdir(parents=True, exist_ok=True)
	workbook.save(output_file)


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Genera un Excel con el concepto y el total por anio, dejando vacio cuando falte el concepto."
	)
	parser.add_argument(
		"--input-dir",
		default=str(PROJECT_ROOT / "DATA" / "SIHO"),
		help="Carpeta SIHO (estructura de SJDC_SIHO) con Formularios_Anuales/<anio>/Produccion/.",
	)
	parser.add_argument(
		"--output-file",
		default=str(PROJECT_ROOT / "SALIDAS" / "produccion" / "tabla_conceptos_por_anio.xlsx"),
		help="Ruta del archivo Excel de salida.",
	)
	parser.add_argument(
		"--concept-column",
		default="concepto",
		help="Columna que contiene el concepto.",
	)
	parser.add_argument(
		"--value-column",
		default="total",
		help="Columna numerica a consolidar por anio.",
	)
	parser.add_argument(
		"--sheet-name",
		default="tabla",
		help="Nombre de la hoja Excel.",
	)
	return parser.parse_args()


def main() -> None:
	args = parse_args()
	input_dir = Path(args.input_dir).resolve()
	output_file = Path(args.output_file).resolve()

	yearly_data = load_yearly_totals(
		input_dir=input_dir,
		concept_column=args.concept_column,
		value_column=args.value_column,
	)
	rows = build_excel_rows(yearly_data)

	if not rows:
		print("No se pudo generar la tabla porque no hubo conceptos validos.")
		return

	write_excel(output_file=output_file, sheet_name=args.sheet_name, rows=rows)

	print(f"Excel generado: {output_file}")
	print(f"Conceptos unicos: {len(rows) - 1}")
	print(f"Anios incluidos: {', '.join(rows[0][1:])}")


if __name__ == "__main__":
	main()