from __future__ import annotations

from pathlib import Path
from typing import Iterable, Dict

from docx import Document


class TemplateProcessor:
    def __init__(self, template_path: Path, placeholders: Dict[str, str]):
        self.template_path = template_path
        self.placeholders = placeholders  # dict: {{NOMBRE}} -> valor
        self.found_placeholders: set[str] = set()

    def render(self) -> Document:
        doc = Document(str(self.template_path))
        # procesar parrafos, tablas y encabezados/pies
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph)
        for section in doc.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    self._replace_in_paragraph(paragraph)
        return doc

    def _replace_in_paragraph(self, paragraph) -> None:
        if not paragraph.runs:
            return
        full_text = "".join(run.text for run in paragraph.runs)
        changed = False
        for ph, val in self.placeholders.items():
            if ph in full_text:
                full_text = full_text.replace(ph, val)
                self.found_placeholders.add(ph)
                changed = True
        if changed:
            # reasignar texto colapsando en un run para simplicidad
            first = paragraph.runs[0]
            first.text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""


__all__ = ["TemplateProcessor"]
